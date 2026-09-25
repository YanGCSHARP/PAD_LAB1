"""Отчёт .docx из experiments/REPORT.md (+ титульный лист и архитектура).

    python experiments/make_figures.py      # SVG
    python experiments/export_png.py        # PNG (через браузер)
    python experiments/build_docx.py        → report/LR1_RAG_Dota2.docx
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
LAB = ROOT.parent
OUT = LAB / "report" / "LR1_RAG_Dota2.docx"
PNG = ROOT / "figures" / "png"

ARCH = """Источники: dota2.com /datafeed (JSON, RU)  +  Dota 2 Fandom (MediaWiki API, EN)
        │
        ▼
Grabber: листинг → проверка версии (revid / патч) → скачивание → raw → парсер
         → Document(секции, metadata) → SHA-256 → манифест SQLite
        │
        ▼
Preprocessing: очистка HTML и вики-мусора → нормализация → chunking
               (fixed / overlap / paragraph / section) + metadata
        │
        ▼
Embeddings (bge-m3, кэш векторов)  →  Vector Store: Qdrant (Docker)
        │
User Query ──► Retriever: hybrid = dense + BM25 (RRF), фильтр по metadata → top-10
        │
        ▼
Filters: длина, дедупликация, порог  →  Reranker: bge-reranker-v2-m3 → top-5
        │
        ▼
LLM (Groq, gpt-oss-120b): системный промпт + контекст [1..5] + вопрос
        │
        ▼
Ответ со ссылками [n] + источники  →  FastAPI  →  Blazor WebAssembly UI"""

INTRO = [
    ("h1", "Цель и постановка задачи"),
    ("p", "Цель работы — построить собственную RAG-систему (Retrieval-Augmented Generation) и "
          "экспериментально исследовать, как архитектурные решения влияют на качество ответов."),
    ("p", "Предметная область — компьютерная игра Dota 2: герои, способности, предметы, игровые "
          "механики и патчи. Система отвечает на вопросы на русском языке только по собранной базе "
          "знаний, указывает источники и сообщает, если информации в базе нет."),
    ("h1", "Архитектура"),
    ("code", ARCH),
    ("p", "Стек: Python (сбор данных, pipeline, эксперименты), Qdrant в Docker (векторная БД), "
          "HuggingFace transformers (embeddings и reranker), Groq и Gemini API (LLM и судья), "
          "FastAPI + Blazor WebAssembly на .NET 10 (веб-интерфейс). Исходный код, инструкции по "
          "запуску и описание модулей — в README.md репозитория."),
]

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text: str, size: float | None = None) -> None:
    text = text.replace("\\*", "*")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("["):
            r = par.add_run(re.match(r"\[([^\]]+)\]", part).group(1))
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[2:]                       # rows[1] — разделитель |---|
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + body):
        for j, val in enumerate(row[:len(header)]):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, val.strip(), size=9)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                shade(cell, "E8EEF7")
    doc.add_paragraph()


def add_code(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(text.split("\n")):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        if i < len(text.split("\n")) - 1:
            r.add_break()


def add_image(doc, svg_rel: str, caption: str) -> None:
    png = PNG / (Path(svg_rel).stem + ".png")
    if not png.exists():
        doc.add_paragraph(f"[нет изображения: {png.name}]")
        return
    doc.add_picture(str(png), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(9)


def title_page(doc) -> None:
    for text, size, bold, space in [
        ("[Наименование учебного заведения]", 12, False, 0), ("[Факультет, кафедра]", 12, False, 120),
        ("ЛАБОРАТОРНАЯ РАБОТА № 1", 16, True, 12), ("Разработка собственной RAG-системы", 14, True, 6),
        ("Предметная область: Dota 2 (официальный сайт dota2.com и Dota 2 Wiki)", 12, False, 150),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
    for text in ("Выполнил: [ФИО, группа]", "Проверил: [ФИО преподавателя]"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph("2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.add_run().add_break(WD_BREAK.PAGE)


def render_markdown(doc, md: str) -> None:
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            add_code(doc, "\n".join(lines[i + 1:j]))
            i = j + 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if m:
            add_image(doc, m.group(2), m.group(1))
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\s*- ", line):
            text = re.sub(r"^\s*- ", "", line)
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and not lines[i + 1].strip().startswith("- "):
                i += 1
                text += " " + lines[i].strip()
            add_runs(doc.add_paragraph(style="List Bullet"), text)
        elif line.strip() in ("", "---"):
            pass
        else:
            text = line
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||!\[|```|\s*- |---|\\\*)", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            add_runs(doc.add_paragraph(), text)
        i += 1


def main() -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    for name in ("Heading 1", "Heading 2", "Title"):
        doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    title_page(doc)
    for kind, val in INTRO:
        if kind == "h1":
            doc.add_heading(val, level=1)
        elif kind == "code":
            add_code(doc, val)
        else:
            add_runs(doc.add_paragraph(), val)

    md = (ROOT / "REPORT.md").read_text(encoding="utf-8")
    md = md.split("\n", 1)[1]                             # заголовок отчёта заменён титульным листом
    render_markdown(doc, md)

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
